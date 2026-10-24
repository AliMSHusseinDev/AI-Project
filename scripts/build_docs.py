"""Build the student report as PDF and editable Word from the Markdown source.

Run with the optional requirements-docs.txt installed. Figures are generated
by scripts/build_figures.py first. No online document service is used.
"""

import html
import re
import textwrap
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate,
                               Spacer, Table, TableStyle, XPreformatted)

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
SOURCES = ['REPORT.md', 'ASSESSMENT.md', 'TEAM.md', 'DISCUSSION.md']


def plain(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', text)
    return text.replace('**', '').replace('`', '')


def inline(text):
    text = html.escape(text)
    text = re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)', r'<link href="\2" color="#176f66">\1</link>', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', text)
    text = re.sub(r'`([^`]+)`', r'<font name="Courier">\1</font>', text)
    return text


def blocks(path):
    lines = path.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i])
                i += 1
            yield 'code', '\n'.join(code)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [cell.strip() for cell in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', cell) for cell in row):
                    rows.append(row)
                i += 1
            yield 'table', rows
            continue
        elif match := re.match(r'^(#{1,6}) (.*)', line):
            yield 'heading', (len(match[1]), match[2])
        elif match := re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line):
            image = (path.parent / match[2]).resolve()
            if image.suffix == '.svg':
                image = image.with_suffix('.png')
            yield 'image', (match[1], image)
        elif re.match(r'^(\d+\. |[-*] )', line):
            yield 'bullet', line
        else:
            paragraph = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||!\[|```|\d+\. |[-*] )', lines[i].strip()):
                paragraph.append(lines[i].strip())
                i += 1
            yield 'paragraph', ' '.join(paragraph)
            continue
        i += 1


def main():
    fontdir = Path('/usr/share/fonts/truetype/dejavu')
    regular, bold = 'Helvetica', 'Helvetica-Bold'
    if (fontdir / 'DejaVuSans.ttf').exists():
        pdfmetrics.registerFont(TTFont('Student', str(fontdir / 'DejaVuSans.ttf')))
        pdfmetrics.registerFont(TTFont('Student-Bold', str(fontdir / 'DejaVuSans-Bold.ttf')))
        pdfmetrics.registerFontFamily('Student', normal='Student', bold='Student-Bold')
        regular, bold = 'Student', 'Student-Bold'
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('StudentBody', fontName=regular, fontSize=9.4, leading=14, spaceAfter=8))
    styles.add(ParagraphStyle('StudentCell', fontName=regular, fontSize=7.4, leading=10))
    styles.add(ParagraphStyle('StudentCaption', fontName=regular, fontSize=8, leading=11,
                             textColor=colors.HexColor('#536b7b'), spaceAfter=12))
    styles.add(ParagraphStyle('StudentCode', fontName='Courier', fontSize=7.4, leading=10,
                             backColor=colors.HexColor('#f1f5f7'), borderPadding=8, spaceAfter=10))
    for name, size, leading in [('Title', 21, 27), ('Heading1', 16, 21), ('Heading2', 12.5, 17), ('Heading3', 10.5, 15)]:
        styles[name].fontName = bold
        styles[name].fontSize = size
        styles[name].leading = leading
        styles[name].textColor = colors.HexColor('#173c4d')
        styles[name].alignment = TA_LEFT
    story = []
    word = Document()
    word.styles['Normal'].font.name = 'Calibri'
    word.styles['Normal'].font.size = Pt(10)
    for section in word.sections:
        section.top_margin = section.bottom_margin = Inches(.7)
        section.left_margin = section.right_margin = Inches(.7)

    for doc_number, filename in enumerate(SOURCES):
        if doc_number:
            story.append(PageBreak())
            word.add_page_break()
        for kind, data in blocks(DOCS / filename):
            if kind == 'heading':
                level, text = data
                style = 'Title' if level == 1 else f'Heading{min(level-1, 3)}'
                story.append(Paragraph(inline(text), styles[style]))
                word.add_heading(plain(text), level=0 if level == 1 else min(level-1, 3))
            elif kind in ('paragraph', 'bullet'):
                story.append(Paragraph(inline(data), styles['StudentBody']))
                word.add_paragraph(plain(data))
            elif kind == 'code':
                wrapped = '\n'.join(textwrap.fill(line, width=91, subsequent_indent='    ',
                                                   replace_whitespace=False) if len(line)>91 else line
                                    for line in data.splitlines())
                story.append(XPreformatted(html.escape(wrapped), styles['StudentCode']))
                paragraph = word.add_paragraph()
                run = paragraph.add_run(data)
                run.font.name, run.font.size = 'Courier New', Pt(8)
            elif kind == 'table':
                cols = len(data[0])
                cells = [[Paragraph(inline(cell), styles['StudentCell']) for cell in row] for row in data]
                table = Table(cells, colWidths=[(A4[0]-88)/cols]*cols, repeatRows=1, hAlign='LEFT')
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e3f1ed')),
                    ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#ccd9df')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
                story.extend([table, Spacer(1, 11)])
                wt = word.add_table(rows=1, cols=cols)
                wt.style = 'Table Grid'
                for i, cell in enumerate(data[0]):
                    wt.rows[0].cells[i].text = plain(cell)
                for row in data[1:]:
                    cells = wt.add_row().cells
                    for i, cell in enumerate(row):
                        cells[i].text = plain(cell)
                word.add_paragraph()
            elif kind == 'image':
                caption, path = data
                if not path.is_file():
                    raise FileNotFoundError(path)
                with PILImage.open(path) as img:
                    width, height = img.size
                factor = min((A4[0]-88)/width, 350/height)
                story.append(Image(str(path), width=width*factor, height=height*factor))
                story.append(Paragraph(html.escape(caption), styles['StudentCaption']))
                word.add_picture(str(path), width=Inches(6.4))
                word.add_paragraph(caption, style='Caption')

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.HexColor('#536b7b'))
        canvas.drawString(44, 26, 'AI 310 / CS361 | Wireless Sensor Placement | Fall 2026')
        canvas.drawRightString(A4[0]-44, 26, str(doc.page))
        canvas.restoreState()

    pdf = SimpleDocTemplate(str(DOCS / 'Project_Report.pdf'), pagesize=A4,
                            rightMargin=44, leftMargin=44, topMargin=42, bottomMargin=43,
                            title='Wireless Sensor Network Placement using GA and MA',
                            author='Student team (names to be completed; AI assistance disclosed)')
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    word.core_properties.title = 'Wireless Sensor Network Placement using GA and MA'
    word.core_properties.author = 'Student team (names to be completed)'
    word.core_properties.comments = 'AI assistance disclosed. Proposed roles are not actual contributions.'
    word.save(DOCS / 'Project_Report.docx')
    from pypdf import PdfReader
    reader = PdfReader(DOCS / 'Project_Report.pdf')
    text = '\n'.join(page.extract_text() or '' for page in reader.pages)
    for required in ['Literature review', '3,000', '13.60', 'Academic', 'TEAM', 'Member']:
        # Case-insensitive check because document headings use sentence case.
        if required.lower() not in text.lower():
            raise RuntimeError(f'PDF is missing expected content: {required}')
    print(f'Built checked PDF ({len(reader.pages)} pages) and editable Word report.')


if __name__ == '__main__':
    main()
